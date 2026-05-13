import os
import re
import sys
import json
import time
import csv
import argparse
import threading
from pathlib import Path
from typing import Any, Dict, List, Optional, Iterable

import requests
from dotenv import load_dotenv, find_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

try:
    from deep_translator import GoogleTranslator
except ImportError:
    GoogleTranslator = None

try:
    from tqdm import tqdm
except ImportError:
    # Fallback if tqdm is not installed
    def tqdm(iterable, **kwargs):
        return iterable

# Importing these at the top now that we've cleaned up db_cache
from db_cache import (
    get_cached_part, save_cached_part, 
    get_translation, save_translation,
    is_docs_category, is_labeling_category, is_ip_brand
)

API_PARTNUMBER = "https://api.mouser.com/api/v1.0/search/partnumber"
API_KEYWORD    = "https://api.mouser.com/api/v1.0/search/keyword"

load_dotenv(find_dotenv(), override=False)

_PROXY_URL = os.getenv("TELEGRAM_PROXY")
_PROXIES = {"http": _PROXY_URL, "https": _PROXY_URL} if _PROXY_URL else None

def _translate_to_ru(text: Optional[str]) -> Optional[str]:
    if not text or not text.strip() or GoogleTranslator is None:
        return text
    text_clean = text.strip()
    
    # 1. Check persistent DB cache
    cached = get_translation(text_clean)
    if cached:
        return cached
    
    # 2. Use API if not found
    try:
        translated = GoogleTranslator(source='auto', target='ru').translate(text_clean)
        if translated:
            save_translation(text_clean, translated)
        return translated
    except Exception as e:
        print(f"WARNING: Translation failed for '{text_clean}': {e}", file=sys.stdout)
        return text

_api_keys = []
_current_key_idx = 0
_key_lock = threading.Lock()

def _init_api_keys():
    global _api_keys
    with _key_lock:
        if _api_keys:
            return
        
        keys = []
        multi_keys = os.getenv("MOUSER_API_KEYS")
        if multi_keys:
            keys.extend([k.strip() for k in multi_keys.split(",") if k.strip()])
            
        for i in range(1, 20):
            k = os.getenv(f"MOUSER_API_KEY_{i}")
            if k and k.strip() not in keys:
                keys.append(k.strip())
                
        single_key = os.getenv("MOUSER_API_KEY")
        if single_key and single_key.strip() not in keys:
            keys.append(single_key.strip())
            
        if not keys:
            print("ERROR: No Mouser API keys found. Set MOUSER_API_KEYS (comma separated) or MOUSER_API_KEY in .env", file=sys.stderr)
            sys.exit(2)
            
        _api_keys = keys
        print(f"Loaded {len(_api_keys)} API keys for rotation.", file=sys.stdout)

def _get_api_key() -> str:
    global _current_key_idx
    _init_api_keys()
    
    with _key_lock:
        key = _api_keys[_current_key_idx]
        _current_key_idx = (_current_key_idx + 1) % len(_api_keys)
        return key

_last_request_time = 0.0
_requests_today = 0
_rate_limit_lock = threading.Lock()

def _enforce_rate_limit():
    global _last_request_time, _requests_today
    
    with _rate_limit_lock:
        now = time.time()
        elapsed = now - _last_request_time
        delay_between_requests = 2.1 / max(1, len(_api_keys))
        
        if elapsed < delay_between_requests:
            sleep_time = delay_between_requests - elapsed
            _last_request_time = now + sleep_time
        else:
            sleep_time = 0.0
            _last_request_time = now
            
        _requests_today += 1
        
    if sleep_time > 0:
        time.sleep(sleep_time)

def _post_json(url: str, params: Dict[str, str], payload: Dict[str, Any],
               max_retries: int = 3, timeout: int = 20) -> Dict[str, Any]:
    headers = {"Content-Type": "application/json", "Accept": "application/json"}
    attempt = 0
    while True:
        attempt += 1
        _enforce_rate_limit()
        try:
            resp = requests.post(
                url, 
                params=params, 
                json=payload, 
                headers=headers, 
                timeout=timeout,
                proxies=_PROXIES
            )
        except requests.RequestException as e:
            if attempt <= max_retries:
                delay = 2 ** attempt
                print(f"Network error. Retrying in {delay}s...", file=sys.stderr)
                time.sleep(delay)
                continue
            raise RuntimeError(f"Network error: {e}") from e

        if resp.status_code == 429:
            if attempt <= max_retries:
                retry_after_str = resp.headers.get("Retry-After", "")
                retry_after = int(retry_after_str) if retry_after_str.isdigit() else 0
                delay = max(retry_after, 2 ** attempt)
                print(f"Rate limit hit (429). Retrying in {delay}s...", file=sys.stderr)
                time.sleep(delay)
                continue
            else:
                raise RuntimeError("Max retries exceeded for 429 Rate Limit.")

        if resp.status_code >= 400:
            raise RuntimeError(f"Mouser API error {resp.status_code}: {resp.text[:500]}")

        try:
            return resp.json()
        except json.JSONDecodeError:
            raise RuntimeError(f"Invalid JSON: {resp.text[:200]}")

def fetch_by_partnumber(query: str, api_key: str, max_retries: int, timeout: int) -> Dict[str, Any]:
    clean_query = query.replace(" ", "").replace("+", "")
    params   = {"apiKey": api_key}
    payload1 = {"SearchByPartNumberRequest": {"MouserPartNumber": clean_query, "records": 50}}
    data = _post_json(API_PARTNUMBER, params, payload1, max_retries, timeout)
    if _has_parts(data):
        return data
    payload2 = {"SearchByPartNumberRequest": {"mouserPartNumber": clean_query, "records": 50}}
    return _post_json(API_PARTNUMBER, params, payload2, max_retries, timeout)

def fetch_by_keyword(query: str, api_key: str, max_retries: int, timeout: int) -> Dict[str, Any]:
    clean_query = query.replace(" ", "").replace("+", "")
    params  = {"apiKey": api_key}
    payload = {"SearchByKeywordRequest": {"keyword": clean_query, "records": 50}}
    return _post_json(API_KEYWORD, params, payload, max_retries, timeout)

def _has_parts(api_response: Dict[str, Any]) -> bool:
    sr = api_response.get("SearchResults") or api_response.get("SearchByPartNumberResponse") or {}
    parts = sr.get("Parts") or []
    return len(parts) > 0

def extract_first_part(api_response: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    sr = api_response.get("SearchResults") or api_response.get("SearchByPartNumberResponse") or {}
    parts = sr.get("Parts") or []
    return parts[0] if parts else None

_AVAIL_RE = re.compile(r"([\d,]+)")
_PN_ALLOWED_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._/+\- ]*$")

def _is_valid_pn(q: str) -> bool:
    q = q.strip()
    if not q:
        return False
    if q in {"\\", "/"}:
        return False
    return bool(_PN_ALLOWED_RE.match(q))

def _parse_stock(availability: Optional[str]) -> Optional[int]:
    if not availability:
        return None
    m = _AVAIL_RE.search(availability)
    if m:
        try:
            return int(m.group(1).replace(",", ""))
        except ValueError:
            pass
    if "on order" in availability.lower():
        return 0
    return None

def _first_unit_price(price_breaks: Optional[List[Dict[str, Any]]]) -> Optional[str]:
    if not price_breaks:
        return None
    for pb in price_breaks:
        try:
            if int(pb.get("Quantity", 0)) == 1:
                p = pb.get("Price")
                return (str(p).strip() if p is not None else None)
        except Exception:
            continue
    p = price_breaks[0].get("Price") if price_breaks else None
    return (str(p).strip() if p is not None else None)

def transform_strict(part: Dict[str, Any], query: str = "") -> Dict[str, Any]:
    mouser_pn    = part.get("MouserPartNumber")
    mfr_pn       = part.get("ManufacturerPartNumber")
    manufacturer = part.get("Manufacturer")
    category     = part.get("Category") or part.get("CategoryName") or part.get("ProductLine") or None
    description  = part.get("Description") or part.get("ProductDescription") or None
    image_url    = part.get("ImagePath") or None
    datasheet    = part.get("DataSheetUrl") or None

    category_ru = _translate_to_ru(category) if category else None
    description_ru = _translate_to_ru(description) if description else None

    res = {}
    if query:
        res["Запрошенный партномер"] = query
    res["Партномер Mouser"] = mouser_pn
    res["Партномер производителя"] = mfr_pn
    res["Производитель"] = manufacturer
    res["Категория (EN)"] = category
    res["Категория"] = category_ru
    res["Описание"] = description_ru
    res["Ссылка на фото"] = image_url
    res["Даташит"] = datasheet

    compliance = part.get("ProductCompliance") or []
    if isinstance(compliance, list):
        for item in compliance:
            if isinstance(item, dict):
                c_name = item.get("ComplianceName")
                c_val = item.get("ComplianceValue")
                if c_name:
                    res[f"Compliance код: {c_name}"] = c_val
    elif isinstance(compliance, dict):
        for k, v in compliance.items():
            res[f"Compliance код: {k}"] = v
    return res

def write_csv(rows: List[Dict[str, Any]], out_path: str) -> None:
    if not rows: return
    fields = []
    for r in rows:
        for k in r.keys():
            if k not in fields:
                fields.append(k)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow(r)

def print_table(rows: List[Dict[str, Any]], display_cols: Optional[List[str]] = None) -> None:
    if not rows: return
    cols = display_cols or []
    if not cols:
        for r in rows:
            for k in r.keys():
                if k not in cols:
                    cols.append(k)
    widths = {c: max(len(c), max((len(str(r.get(c) or "")) for r in rows), default=0)) for c in cols}
    header = " | ".join(c.ljust(widths[c]) for c in cols)
    sep = "-+-".join("-" * widths[c] for c in cols)
    print(header); print(sep)
    for r in rows:
        print(" | ".join(str(r.get(c) or "").ljust(widths[c]) for c in cols))

def _iter_input(parts_from_cli: List[str], input_file: Optional[str]) -> Iterable[str]:
    if input_file:
        p = Path(input_file)
        if not p.exists():
            raise FileNotFoundError(f"--input file not found: {input_file}")
        if p.suffix.lower() == ".csv":
            with p.open("r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if not line: continue
                    yield line.split(",")[0].strip()
        else:
            with p.open("r", encoding="utf-8") as f:
                for line in f:
                    s = line.strip()
                    if s: yield s
    for x in parts_from_cli:
        yield x

def _to_number_or_text(value):
    if value is None: return None
    s = str(value).strip()
    if s.isdigit():
        try: return int(s)
        except: return s
    try:
        s2 = s[1:] if s[0] in ("$", "€", "£") else s
        return float(s2.replace(",", ""))
    except: return s

def write_xlsx(rows: List[Dict[str, Any]], out_path: str) -> None:
    if not rows: return
    wb = Workbook()
    ws = wb.active
    ws.title = "Mouser"
    cols = []
    for r in rows:
        for k in r.keys():
            if k not in cols: cols.append(k)
    header_font = Font(bold=True)
    ws.append(cols)
    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    pale_yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    light_red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    light_blue_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    ip_font = Font(bold=True, color="7030A0")

    # Добавляем легенду сверху
    ws.insert_rows(1)
    ws.cell(row=1, column=1, value="ЛЕГЕНДА:").font = Font(bold=True)
    
    cell_doc = ws.cell(row=1, column=2, value="Нужны разр. документы")
    cell_doc.fill = pale_yellow_fill
    
    cell_lbl = ws.cell(row=1, column=3, value="Требуется маркировка")
    cell_lbl.fill = light_red_fill
    
    cell_both = ws.cell(row=1, column=4, value="Документы + Маркировка")
    cell_both.fill = light_blue_fill
    
    cell_ip = ws.cell(row=1, column=5, value="Обнаружено в ТРОИС")
    cell_ip.font = ip_font

    for r in rows:
        out_row = [_to_number_or_text(r.get(c)) for c in cols]
        ws.append(out_row)
        category_en = r.get("Категория (EN)")
        category_ru = r.get("Категория")
        is_labeled = False
        is_doc_req = False
        for c_val in [category_en, category_ru]:
            if c_val and c_val != "-":
                if is_labeling_category(c_val): is_labeled = True
                if is_docs_category(c_val): is_doc_req = True
        
        fill_to_apply = None
        if is_labeled and is_doc_req: fill_to_apply = light_blue_fill
        elif is_labeled: fill_to_apply = light_red_fill
        elif is_doc_req: fill_to_apply = pale_yellow_fill
        
        mfr_api = r.get("Производитель")
        brand_excel = r.get("brand_from_excel")
        font_to_apply = ip_font if (brand_excel and is_ip_brand(brand_excel)) or (mfr_api and is_ip_brand(mfr_api)) else None

        if fill_to_apply or font_to_apply:
            for cell in ws[ws.max_row]:
                if fill_to_apply: cell.fill = fill_to_apply
                if font_to_apply: cell.font = font_to_apply
    ws.freeze_panes = "A2"
    for col_idx, col_name in enumerate(cols, start=1):
        max_len = len(str(col_name))
        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val: max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)
    wb.save(out_path)
    print(f"Wrote XLSX -> {out_path}")

def enrich_xlsx(input_path: str, out_path: str, results_map: Dict[str, Dict[str, Any]]) -> None:
    """Обогащает существующий Excel файл данными из Mouser."""
    try:
        wb = load_workbook(input_path, data_only=True)
        ws = wb.active
    except Exception as e:
        print(f"Error loading workbook for enrichment: {e}")
        return

    # 1. Ищем колонку с партномером (Part no.)
    part_col_idx = None
    start_row = None
    pn_col_names = ("partno", "partnumber", "partnum", "pn", "p/n")

    for row_idx, row in enumerate(ws.iter_rows(max_row=50), start=1):
        for col_idx, cell in enumerate(row, start=1):
            if cell.value and isinstance(cell.value, str):
                val_clean = cell.value.strip().lower().replace(" ", "").replace(".", "")
                if val_clean in pn_col_names:
                    part_col_idx = col_idx
                    start_row = row_idx + 1
                    break
        if part_col_idx: break

    if not part_col_idx:
        print("Could not find PN column for enrichment.")
        return

    # 2. Определяем, куда писать данные (первая пустая колонка после используемого диапазона)
    target_start_col = ws.max_column + 1
    
    # Определяем заголовки для добавления
    new_cols = [
        "Партномер Mouser", "Производитель", "Категория", 
        "Описание", "Ссылка на фото", "Даташит", "Статус ТРОИС/Маркировка"
    ]
    
    header_font = Font(bold=True)
    for i, col_name in enumerate(new_cols):
        cell = ws.cell(row=start_row - 1, column=target_start_col + i)
        cell.value = col_name
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Стили
    pale_yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    light_red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    light_blue_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    ip_font = Font(bold=True, color="7030A0")

    # Добавляем легенду сверху
    ws.insert_rows(1)
    if start_row: start_row += 1
    
    ws.cell(row=1, column=1, value="ЛЕГЕНДА:").font = Font(bold=True)
    ws.cell(row=1, column=2, value="Нужны разр. документы").fill = pale_yellow_fill
    ws.cell(row=1, column=3, value="Требуется маркировка").fill = light_red_fill
    ws.cell(row=1, column=4, value="Документы + Маркировка").fill = light_blue_fill
    ws.cell(row=1, column=5, value="Обнаружено в ТРОИС").font = ip_font

    # 3. Заполняем данными
    for row_idx in range(start_row, ws.max_row + 1):
        pn_val = ws.cell(row=row_idx, column=part_col_idx).value
        if not pn_val: continue
        
        pn_str = str(pn_val).strip().lower()
        # Ищем в мапе (ключи там в нижнем регистре)
        res = results_map.get(pn_str)
        if not res: continue

        # Заполняем ячейки
        ws.cell(row=row_idx, column=target_start_col + 0).value = res.get("Партномер Mouser", "-")
        ws.cell(row=row_idx, column=target_start_col + 1).value = res.get("Производитель", "-")
        ws.cell(row=row_idx, column=target_start_col + 2).value = res.get("Категория", "-")
        ws.cell(row=row_idx, column=target_start_col + 3).value = res.get("Описание", "-")
        ws.cell(row=row_idx, column=target_start_col + 4).value = res.get("Ссылка на фото", "-")
        ws.cell(row=row_idx, column=target_start_col + 5).value = res.get("Даташит", "-")
        
        # Статус (ТРОИС / Маркировка / Доки)
        warnings = []
        cat_en = res.get("Категория (EN)")
        cat_ru = res.get("Категория")
        
        is_labeled = False
        is_doc_req = False
        
        # Проверяем оба варианта названия категории
        for c_val in [cat_en, cat_ru]:
            if c_val and c_val != "-":
                if is_labeling_category(c_val): is_labeled = True
                if is_docs_category(c_val): is_doc_req = True
        
        if is_labeled: warnings.append("Маркировка")
        if is_doc_req: warnings.append("Документы")
        
        mfr = res.get("Производитель")
        brand_excel = res.get("brand_from_excel")
        is_ip = False
        
        # Сверка с ТРОИС: приоритет данным из инвойса
        brand_to_check = brand_excel if brand_excel else mfr
        if brand_to_check and brand_to_check != "-" and is_ip_brand(brand_to_check):
            is_ip = True
            warnings.append("ТРОИС")

        status_text = ", ".join(warnings) if warnings else "-"
        ws.cell(row=row_idx, column=target_start_col + 6).value = status_text

        # Применяем стили
        fill_to_apply = None
        if is_labeled and is_doc_req:
            fill_to_apply = light_blue_fill
        elif is_labeled:
            fill_to_apply = light_red_fill
        elif is_doc_req:
            fill_to_apply = pale_yellow_fill
        
        if fill_to_apply:
            for i in range(1, target_start_col + len(new_cols)):
                ws.cell(row=row_idx, column=i).fill = fill_to_apply
        
        if is_ip:
            for i in range(1, target_start_col + len(new_cols)):
                ws.cell(row=row_idx, column=i).font = ip_font

    # Авто-ширина новых колонок
    for i in range(len(new_cols)):
        col_idx = target_start_col + i
        max_len = len(new_cols[i])
        for row_idx in range(start_row, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val: max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

    wb.save(out_path)

def main():
    parser = argparse.ArgumentParser(description="Mouser Search CLI: Mouser PN or Manufacturer PN (keyword).")
    parser.add_argument("part_numbers", nargs="*", help="PNs (Mouser or Manufacturer)")
    parser.add_argument("--input", help="Path to file with PNs")
    parser.add_argument("--format", choices=["table", "csv", "json", "xlsx"], default="table")
    parser.add_argument("--out", help="Output file path")
    parser.add_argument("--raw", action="store_true", help="Print raw JSON")
    parser.add_argument("--save-raw", help="Directory to save raw JSON")
    parser.add_argument("--timeout", type=int, default=20, help="HTTP timeout")
    parser.add_argument("--retries", type=int, default=3, help="Max retries")
    args = parser.parse_args()

    api_key = _get_api_key()
    q_iter = list(_iter_input(args.part_numbers, args.input))
    if not q_iter:
        print("Provide at least one PN via args or --input file.", file=sys.stderr)
        sys.exit(2)

    results: List[Dict[str, Any]] = []
    save_dir = Path(args.save_raw) if args.save_raw else None
    if save_dir: save_dir.mkdir(parents=True, exist_ok=True)

    for q in tqdm(q_iter, desc="Processing PNs", disable=len(q_iter) < 5):
        cached = get_cached_part(q)
        if cached:
            results.append(cached)
            continue
        if not _is_valid_pn(q):
            print(f"WARNING: skipped invalid query token: {q!r}", file=sys.stderr)
            continue
        data_pn = fetch_by_partnumber(q, api_key, args.retries, args.timeout)
        part = extract_first_part(data_pn)
        used_keyword = False
        if not part:
            data_kw = fetch_by_keyword(q, api_key, args.retries, args.timeout)
            part = extract_first_part(data_kw)
            used_keyword = True
        if args.raw:
            print(json.dumps(data_kw if used_keyword else data_pn, ensure_ascii=False, indent=2))
        if save_dir:
            raw_path = save_dir / f"{q.replace('/', '_')}.json"
            with raw_path.open("w", encoding="utf-8") as f:
                json.dump(data_kw if used_keyword else data_pn, f, ensure_ascii=False, indent=2)
        if not part:
            empty_res = {"Запрошенный партномер": q, "Партномер Mouser": "-", "Партномер производителя": "-", "Производитель": "-", "Категория (EN)": "-", "Категория": "-", "Описание": "Ничего не найдено", "Ссылка на фото": "-", "Даташит": "-"}
            save_cached_part(q, empty_res)
            results.append(empty_res)
            continue
        row = transform_strict(part, q)
        for k, v in row.items():
            if v is None: row[k] = "-"
        save_cached_part(q, row)
        results.append(row)

    if not results: return
    if args.format == "json":
        output = json.dumps(results, ensure_ascii=False, indent=2)
        if args.out:
            with open(args.out, "w", encoding="utf-8") as f: f.write(output)
            print(f"Wrote JSON -> {args.out}")
        else: print(output)
    elif args.format == "csv":
        out_path = args.out or "mouser_results.csv"
        write_csv(results, out_path)
        print(f"Wrote CSV -> {out_path}")
    elif args.format == "xlsx":
        out_path = args.out or "mouser_results.xlsx"
        write_xlsx(results, out_path)
    else:
        print_table(results, display_cols=["Запрошенный партномер", "Партномер Mouser", "Производитель", "Категория", "Описание"])

if __name__ == "__main__":
    main()



if __name__ == "__main__":
    main()
