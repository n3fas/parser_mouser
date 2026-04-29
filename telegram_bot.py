import os
import sys
import asyncio
import tempfile
import math
import datetime
from pathlib import Path

from dotenv import load_dotenv, find_dotenv

from mouser_logger import setup_bot_logger
logger = setup_bot_logger()

from aiogram import Bot, Dispatcher, F, types
from aiogram.filters import CommandStart, Command
from aiogram.types import Message, FSInputFile, InlineKeyboardMarkup, InlineKeyboardButton, CallbackQuery, LinkPreviewOptions, ReplyKeyboardRemove, BotCommand
from aiogram.enums import ParseMode
from aiogram.client.default import DefaultBotProperties
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from openpyxl import load_workbook, Workbook
from aiogram.dispatcher.middlewares.base import BaseMiddleware
from typing import Callable, Dict, Any, Awaitable, Optional
from aiogram.client.session.aiohttp import AiohttpSession

import aiohttp
try:
    from aiohttp_socks import ProxyConnector
except ImportError:
    ProxyConnector = None

class UpdateState(StatesGroup):
    waiting_for_data = State()

class DocsCategoryState(StatesGroup):
    waiting_for_add = State()
    waiting_for_remove = State()
    waiting_for_search = State()

class LabelingCategoryState(StatesGroup):
    waiting_for_add = State()
    waiting_for_remove = State()
    waiting_for_search = State()

class IPRegistryState(StatesGroup):
    waiting_for_file = State()

class UserManagementState(StatesGroup):
    waiting_for_add_id = State()
    waiting_for_remove_id = State()

# Импортируем готовые функции из нашего проекта
from mouser_cli import _get_api_key, write_xlsx, enrich_xlsx, _translate_to_ru
from mouser_menu import _lookup_one, _split_pns
from db_cache import (
    add_to_history, get_user_history, clear_user_history, 
    get_cache_stats, delete_cached_part, get_all_cached_pns,
    add_docs_category, remove_docs_category, get_docs_categories, get_docs_categories_count, is_docs_category,
    add_labeling_category, remove_labeling_category, get_labeling_categories, get_labeling_categories_count, is_labeling_category,
    clear_ip_brands, add_ip_brand, is_ip_brand, get_ip_brands_count, get_all_ip_brands,
    add_user, remove_user, is_user_allowed, get_all_users
)

load_dotenv(find_dotenv(), override=False)

TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
ADMIN_USER_ID = os.getenv("ADMIN_USER_ID")
PROXY_URL = os.getenv("TELEGRAM_PROXY")
if not TOKEN:
    logger.warning("Переменная TELEGRAM_BOT_TOKEN не задана в .env файле.")

ALLOWED_THREAD_ID = os.getenv("ALLOWED_THREAD_ID")
if ALLOWED_THREAD_ID:
    try: ALLOWED_THREAD_ID = int(ALLOWED_THREAD_ID)
    except ValueError: logger.warning("ALLOWED_THREAD_ID должен быть числом.")

if ADMIN_USER_ID:
    try: ADMIN_USER_ID = int(ADMIN_USER_ID)
    except ValueError: logger.warning("ADMIN_USER_ID должен быть числом.")

dp = Dispatcher()

def get_menu_text():
    return (
        "🤖 <b>Как работает парсер Mouser:</b>\n\n"
        "Отправьте мне:\n"
        "1️⃣ <b>Текст</b> с парт-номерами (через пробел или запятую) - я верну информацию прямо сюда.\n"
        "2️⃣ <b>Excel-файл</b> (ВАЖНО: формат - .xlsx) - я найду в нем столбец 'Part no.' (или 'PART NUMBER'), "
        "соберу все номера и отправлю обратно заполненный Excel-файл со всеми данными, включая перевод."
    )

def get_inline_menu_keyboard(user_id: int):
    kb = [
        [InlineKeyboardButton(text="🔄 Актуализация", callback_data="menu_update")],
        [InlineKeyboardButton(text="🔄 Актуализация ПОЛНАЯ", callback_data="menu_update_all")],
        [InlineKeyboardButton(text="🛡 Реестр ТРОИС (Интеллектуалка)", callback_data="menu_ip_registry")],
        [InlineKeyboardButton(text="📋 Категории для доков", callback_data="menu_docs_cats")],
        [InlineKeyboardButton(text="🏷 Категории для маркировки", callback_data="menu_labeling_cats")],
        [InlineKeyboardButton(text="📊 Статистика базы", callback_data="menu_stats")],
        [InlineKeyboardButton(text="🕒 История запросов", callback_data="menu_history")],
        [InlineKeyboardButton(text="🗑 Очистить историю", callback_data="menu_clear_history")]
    ]
    if ADMIN_USER_ID and user_id == ADMIN_USER_ID:
        kb.insert(2, [InlineKeyboardButton(text="👤 Управление пользователями", callback_data="menu_users")])
    return InlineKeyboardMarkup(inline_keyboard=kb)

def get_docs_cats_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Добавить", callback_data="docs_cat_add")],
        [InlineKeyboardButton(text="➖ Удалить", callback_data="docs_cat_remove")],
        [InlineKeyboardButton(text="🔍 Поиск", callback_data="docs_cat_search")],
        [InlineKeyboardButton(text="📜 Список", callback_data="docs_cat_list")],
        [InlineKeyboardButton(text="🔙 Назад", callback_data="menu_back")]
    ])

def get_labeling_cats_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Добавить", callback_data="labeling_cat_add")],
        [InlineKeyboardButton(text="➖ Удалить", callback_data="labeling_cat_remove")],
        [InlineKeyboardButton(text="🔍 Поиск", callback_data="labeling_cat_search")],
        [InlineKeyboardButton(text="📜 Список", callback_data="labeling_cat_list")],
        [InlineKeyboardButton(text="🔙 Назад", callback_data="menu_back")]
    ])

_user_last_pns = {}

def process_pns_sync(pns, api_key, force_update=False):
    results = []
    for pn in pns:
        row = _lookup_one(pn, api_key, force_update=force_update)
        if row: results.append(row)
    return results

def extract_pns_from_excel_sync(filepath):
    try:
        wb = load_workbook(filepath, data_only=True); ws = wb.active
    except Exception as e: return None, f"Ошибка: {e}"
    part_col_idx = None; brand_col_idx = None; start_row = None
    pn_names = ("partno", "partnumber", "partnum", "pn", "p/n")
    br_names = ("brand", "mfr", "mark", "manufacturer")
    for r_idx, row in enumerate(ws.iter_rows(max_row=50), 1):
        for c_idx, cell in enumerate(row, 1):
            if cell.value and isinstance(cell.value, str):
                v = cell.value.strip().lower().replace(" ", "").replace(".", "")
                if v in pn_names and not part_col_idx: part_col_idx = c_idx
                elif v in br_names and not brand_col_idx: brand_col_idx = c_idx
        if part_col_idx: start_row = r_idx + 1; break
    if not part_col_idx: return None, "NO_COLUMN"
    parts = []
    for r_idx in range(start_row, ws.max_row + 1):
        pn = ws.cell(row=r_idx, column=part_col_idx).value
        if pn is None or str(pn).strip() == "": break
        br = ws.cell(row=r_idx, column=brand_col_idx).value if brand_col_idx else None
        parts.append({"pn": str(pn).strip(), "brand": str(br).strip() if br else None})
    return parts, None

def create_template_excel():
    wb = Workbook(); ws = wb.active; ws.append(["Part no."])
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); wb.save(tmp.name)
    return tmp.name

def get_progress_bar(current, total, length=10):
    p = current / total if total > 0 else 0
    f = int(length * p); bar = "⬛" * f + "⬜" * (length - f)
    return f"{bar} ({int(p * 100)}%)"

def translate_cats_sync(cats):
    res = []
    for c in cats:
        t = _translate_to_ru(c)
        res.append(f"• <code>{c}</code> ({t})" if t and t != c else f"• <code>{c}</code>")
    return res

@dp.message(CommandStart())
async def command_start_handler(message: Message) -> None:
    await message.answer("Привет! Нажмите «Меню» для начала.", reply_markup=ReplyKeyboardRemove())

@dp.message(Command("menu"))
async def handle_menu_command(message: Message) -> None:
    await message.answer(get_menu_text(), reply_markup=get_inline_menu_keyboard(message.from_user.id))

@dp.callback_query(F.data == "menu_docs_cats")
async def callback_menu_docs_cats(callback: CallbackQuery):
    await callback.message.edit_text("📋 <b>Категории для доков</b>\n(бледно-желтый в Excel)", reply_markup=get_docs_cats_keyboard())
    await callback.answer()

@dp.callback_query(F.data == "menu_back")
async def callback_menu_back(callback: CallbackQuery, state: FSMContext):
    await state.clear(); await callback.message.edit_text(get_menu_text(), reply_markup=get_inline_menu_keyboard(callback.from_user.id))
    await callback.answer()

@dp.callback_query(F.data == "menu_users")
async def callback_menu_users(callback: CallbackQuery):
    users = get_all_users()
    text = f"👤 <b>Пользователи:</b>\n\n" + ("\n".join([f"• <code>{u}</code>" for u in users]) if users else "Пусто.")
    kb = [[InlineKeyboardButton(text="➕", callback_data="user_add"), InlineKeyboardButton(text="➖", callback_data="user_remove")], [InlineKeyboardButton(text="🔙", callback_data="menu_back")]]
    await callback.message.edit_text(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=kb))
    await callback.answer()

@dp.callback_query(F.data == "user_add")
async def callback_user_add(callback: CallbackQuery, state: FSMContext):
    await state.set_state(UserManagementState.waiting_for_add_id)
    await callback.message.answer("ID пользователя?")
    await callback.answer()

@dp.callback_query(F.data == "user_remove")
async def callback_user_remove(callback: CallbackQuery, state: FSMContext):
    await state.set_state(UserManagementState.waiting_for_remove_id)
    await callback.message.answer("ID для удаления?")
    await callback.answer()

@dp.message(UserManagementState.waiting_for_add_id)
async def handle_user_add(message: Message, state: FSMContext):
    uid = message.forward_from.id if message.forward_from else (int(message.text) if message.text.isdigit() else None)
    if uid and add_user(uid): await message.answer(f"✅ {uid} добавлен."); await state.clear()
    else: await message.answer("Ошибка.")

@dp.message(UserManagementState.waiting_for_remove_id)
async def handle_user_remove(message: Message, state: FSMContext):
    uid = int(message.text) if message.text.isdigit() else None
    if uid and remove_user(uid): await message.answer(f"✅ {uid} удален."); await state.clear()
    else: await message.answer("Ошибка.")

@dp.callback_query(F.data.startswith("docs_list_") | (F.data == "docs_cat_list"))
async def callback_docs_cat_list(callback: CallbackQuery):
    page = 0 if callback.data == "docs_cat_list" else int(callback.data.split("_")[2])
    total = get_docs_categories_count(); cats = get_docs_categories(limit=20, offset=page*20)
    text = f"📜 <b>Доки ({page+1}):</b>\n\n" + "\n".join(await asyncio.to_thread(translate_cats_sync, cats))
    kb = []
    if page > 0: kb.append(InlineKeyboardButton(text="⬅️", callback_data=f"docs_list_{page-1}"))
    if (page+1)*20 < total: kb.append(InlineKeyboardButton(text="➡️", callback_data=f"docs_list_{page+1}"))
    markup = [kb, [InlineKeyboardButton(text="🔙", callback_data="menu_docs_cats")]]
    await callback.message.edit_text(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=markup))
    await callback.answer()

@dp.message(Command("stats"))
async def handle_stats_command(message: Message):
    count = get_cache_stats()
    text = (
        f"📊 <b>Статистика базы:</b>\n\n"
        f"• Деталей в кэше: {count}\n"
        f"• Разр. документы (категории): {get_docs_categories_count()}\n"
        f"• Маркировка (категории): {get_labeling_categories_count()}\n"
        f"• ТРОИС (брендов): {get_ip_brands_count()}\n"
        f"• Пользователей: {len(get_all_users())}"
    )
    await message.answer(text)

@dp.callback_query(F.data == "menu_stats")
async def callback_menu_stats(callback: CallbackQuery):
    count = get_cache_stats()
    text = (
        f"📊 <b>Статистика базы:</b>\n\n"
        f"• Деталей в кэше: {count}\n"
        f"• Разр. документы (категории): {get_docs_categories_count()}\n"
        f"• Маркировка (категории): {get_labeling_categories_count()}\n"
        f"• ТРОИС (брендов): {get_ip_brands_count()}\n"
        f"• Пользователей: {len(get_all_users())}"
    )
    kb = [[InlineKeyboardButton(text="🔙 Назад", callback_data="menu_back")]]
    await callback.message.edit_text(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=kb))
    await callback.answer()

@dp.message(Command("history"))
async def handle_history_command(message: Message):
    history = get_user_history(message.from_user.id)
    if not history: await message.answer("История пуста."); return
    text = "🕒 <b>Последние 20 запросов:</b>\n\n" + "\n".join([f"• <code>{pn}</code>" for pn in history])
    await message.answer(text)

@dp.callback_query(F.data == "menu_history")
async def callback_menu_history(callback: CallbackQuery):
    history = get_user_history(callback.from_user.id)
    text = "🕒 <b>Последние 20 запросов:</b>\n\n" + ("\n".join([f"• <code>{pn}</code>" for pn in history]) if history else "История пуста.")
    kb = [[InlineKeyboardButton(text="🔙 Назад", callback_data="menu_back")]]
    await callback.message.edit_text(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=kb))
    await callback.answer()

@dp.message(Command("clear_history"))
async def handle_clear_history_command(message: Message):
    clear_user_history(message.from_user.id); await message.answer("✅ История очищена.")

@dp.callback_query(F.data == "menu_clear_history")
async def callback_menu_clear_history(callback: CallbackQuery):
    clear_user_history(callback.from_user.id); await callback.answer("История очищена!")
    await callback_menu_back(callback, None)

@dp.callback_query(F.data == "docs_cat_add")
async def callback_docs_cat_add(callback: CallbackQuery, state: FSMContext):
    await state.set_state(DocsCategoryState.waiting_for_add); await callback.message.answer("Категория (текст или файл)?"); await callback.answer()

@dp.callback_query(F.data == "docs_cat_remove")
async def callback_docs_cat_remove(callback: CallbackQuery, state: FSMContext):
    await state.set_state(DocsCategoryState.waiting_for_remove); await callback.message.answer("Категория для удаления?"); await callback.answer()

@dp.message(DocsCategoryState.waiting_for_add)
async def handle_docs_cat_add(message: Message, state: FSMContext):
    if message.document: return await handle_document(message, message.bot, state)
    lines = [l.strip() for l in message.text.split('\n') if l.strip()]
    added = sum(1 for c in lines if add_docs_category(c))
    await message.answer(f"✅ Добавлено: {added}"); await state.clear()

@dp.message(DocsCategoryState.waiting_for_remove)
async def handle_docs_cat_remove(message: Message, state: FSMContext):
    if remove_docs_category(message.text.strip()): await message.answer("✅ Удалено.")
    else: await message.answer("⚠️ Не найдено."); await state.clear()

@dp.callback_query(F.data == "menu_labeling_cats")
async def callback_menu_labeling_cats(callback: CallbackQuery):
    await callback.message.edit_text("🏷 <b>Категории для маркировки</b>\n(бледно-зеленый в Excel)", reply_markup=get_labeling_cats_keyboard())
    await callback.answer()

@dp.callback_query(F.data.startswith("labeling_list_") | (F.data == "labeling_cat_list"))
async def callback_labeling_cat_list(callback: CallbackQuery):
    page = 0 if callback.data == "labeling_cat_list" else int(callback.data.split("_")[2])
    total = get_labeling_categories_count(); cats = get_labeling_categories(limit=20, offset=page*20)
    text = f"📜 <b>Маркировка ({page+1}):</b>\n\n" + "\n".join(await asyncio.to_thread(translate_cats_sync, cats))
    kb = []
    if page > 0: kb.append(InlineKeyboardButton(text="⬅️", callback_data=f"labeling_list_{page-1}"))
    if (page+1)*20 < total: kb.append(InlineKeyboardButton(text="➡️", callback_data=f"labeling_list_{page+1}"))
    markup = [kb, [InlineKeyboardButton(text="🔙", callback_data="menu_labeling_cats")]]
    await callback.message.edit_text(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=markup))
    await callback.answer()

@dp.callback_query(F.data == "labeling_cat_add")
async def callback_labeling_cat_add(callback: CallbackQuery, state: FSMContext):
    await state.set_state(LabelingCategoryState.waiting_for_add); await callback.message.answer("Категория (текст или файл)?"); await callback.answer()

@dp.callback_query(F.data == "labeling_cat_remove")
async def callback_labeling_cat_remove(callback: CallbackQuery, state: FSMContext):
    await state.set_state(LabelingCategoryState.waiting_for_remove); await callback.message.answer("Категория для удаления?"); await callback.answer()

@dp.message(LabelingCategoryState.waiting_for_add)
async def handle_labeling_cat_add(message: Message, state: FSMContext):
    if message.document: return await handle_document(message, message.bot, state)
    lines = [l.strip() for l in message.text.split('\n') if l.strip()]
    added = sum(1 for c in lines if add_labeling_category(c))
    await message.answer(f"✅ Добавлено: {added}"); await state.clear()

@dp.message(LabelingCategoryState.waiting_for_remove)
async def handle_labeling_cat_remove(message: Message, state: FSMContext):
    if remove_labeling_category(message.text.strip()): await message.answer("✅ Удалено.")
    else: await message.answer("⚠️ Не найдено."); await state.clear()

@dp.callback_query(F.data == "menu_update")
async def callback_menu_update(callback: CallbackQuery, state: FSMContext):
    await state.set_state(UpdateState.waiting_for_data); await callback.message.answer("🔄 Пришлите PN или Excel (обновление)..."); await callback.answer()

@dp.callback_query(F.data == "start_update_all")
async def callback_start_update_all(callback: CallbackQuery):
    pns = get_all_cached_pns()
    if not pns: await callback.answer("Пусто."); return
    await callback.message.edit_text(f"⏳ Обновляю {len(pns)} деталей..."); asyncio.create_task(run_full_update(callback.message, pns)); await callback.answer()

async def run_full_update(message: Message, pns: list):
    try: api_key = _get_api_key()
    except: await message.answer("Ключ не найден."); return
    total = len(pns); processed = 0; last_t = asyncio.get_event_loop().time()
    for pn in pns:
        await asyncio.to_thread(_lookup_one, pn, api_key, force_update=True); processed += 1
        now = asyncio.get_event_loop().time()
        if now - last_t > 3.0 or processed == total:
            try: await message.edit_text(f"⏳ {processed}/{total}\n{get_progress_bar(processed, total)}"); last_t = now
            except: pass
    await message.answer(f"✅ Готово: {total}.")

@dp.callback_query(F.data == "menu_ip_registry")
async def callback_menu_ip_registry(callback: CallbackQuery, state: FSMContext):
    await state.set_state(IPRegistryState.waiting_for_file)
    await callback.message.edit_text(f"🛡 <b>ТРОИС</b> (брендов: {get_ip_brands_count()})\n\nПришлите файл для обновления.", reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="📤 Выгрузить", callback_data="export_ip_brands")], [InlineKeyboardButton(text="🔙", callback_data="menu_back")]]))
    await callback.answer()

@dp.callback_query(F.data == "export_ip_brands")
async def callback_export_ip_brands(callback: CallbackQuery):
    brands = get_all_ip_brands()
    if not brands: await callback.answer("Пусто."); return
    msg = await callback.message.answer("⏳ Формирую..."); tmp_path = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name
    def write(b_list, p):
        wb = Workbook(); ws = wb.active; ws.append(["Brand Name"])
        for b in b_list: ws.append([b])
        wb.save(p)
    await asyncio.to_thread(write, brands, tmp_path)
    await callback.message.answer_document(FSInputFile(tmp_path, filename="ip_brands.xlsx")); os.remove(tmp_path); await msg.delete()

def format_text_result(results):
    lines = []
    for r in results:
        req_pn = r.get("Запрошенный партномер", "-")
        mfr = r.get("Производитель", "-")
        cat_ru = r.get("Категория", "-")
        cat_en = r.get("Категория (EN)")
        desc = r.get("Описание", "-")
        ds = r.get("Даташит"); img = r.get("Ссылка на фото")
        lines.append(f"🔍 <b>{req_pn}</b>\nПроизводитель: {mfr}\nКатегория: {cat_ru}")
        warns = []
        # Проверяем и английское, и русское название категории
        for c_val in [cat_en, cat_ru]:
            if c_val and c_val != "-":
                if is_labeling_category(c_val) and "🟢 Маркировка" not in warns: warns.append("🟢 Маркировка")
                if is_docs_category(c_val) and "🟡 Документы" not in warns: warns.append("🟡 Документы")
        brand_excel = r.get("brand_from_excel")
        if (brand_excel and is_ip_brand(brand_excel)) or (mfr and mfr != "-" and is_ip_brand(mfr)): warns.append("⛔ ТРОИС")
        if warns: lines.append("⚠️ <b>Внимание:</b> " + ", ".join(warns))
        lines.append(f"Описание: {desc}")
        links = []
        if ds and ds != "-": links.append(f'<a href="{ds}">📄 Даташит</a>')
        if img and img != "-": links.append(f'<a href="{img}">🖼 Фото</a>')
        if links: lines.append(" | ".join(links))
        lines.append("-" * 25)
    return "\n".join(lines)

@dp.callback_query(F.data == "docs_cat_search")
async def callback_docs_cat_search(callback: CallbackQuery, state: FSMContext):
    await state.set_state(DocsCategoryState.waiting_for_search); await callback.message.answer("Поиск в категориях ДОКОВ (часть названия):"); await callback.answer()

@dp.callback_query(F.data == "labeling_cat_search")
async def callback_labeling_cat_search(callback: CallbackQuery, state: FSMContext):
    await state.set_state(LabelingCategoryState.waiting_for_search); await callback.message.answer("Поиск в категориях МАРКИРОВКИ (часть названия):"); await callback.answer()

@dp.message(DocsCategoryState.waiting_for_search)
async def handle_docs_cat_search(message: Message, state: FSMContext):
    q = message.text.strip().lower(); res = [c for c in get_docs_categories() if q in c.lower()]
    if not res: await message.answer("Ничего не найдено.")
    else: await message.answer(f"🔍 Найдено ({len(res)}):\n\n" + "\n".join(await asyncio.to_thread(translate_cats_sync, res)))
    await state.clear()

@dp.message(LabelingCategoryState.waiting_for_search)
async def handle_labeling_cat_search(message: Message, state: FSMContext):
    q = message.text.strip().lower(); res = [c for c in get_labeling_categories() if q in c.lower()]
    if not res: await message.answer("Ничего не найдено.")
    else: await message.answer(f"🔍 Найдено ({len(res)}):\n\n" + "\n".join(await asyncio.to_thread(translate_cats_sync, res)))
    await state.clear()

@dp.message(F.text & ~F.text.startswith('/'))
async def handle_text(message: Message, bot: Bot, state: FSMContext) -> None:
    current_state = await state.get_state()
    # Игнорируем, если есть активное состояние (кроме UpdateState)
    if current_state and current_state != UpdateState.waiting_for_data.state: return

    force_update = current_state == UpdateState.waiting_for_data.state
    if force_update: await state.clear()
    if message.chat.type != "private":
        if ALLOWED_THREAD_ID and message.message_thread_id != ALLOWED_THREAD_ID: return
        bot_user = await bot.get_me()
        if not (message.reply_to_message and message.reply_to_message.from_user.id == bot_user.id or f"@{bot_user.username}" in message.text): return
    try: api_key = _get_api_key()
    except: await message.answer("Ключ не найден."); return
    bot_user = await bot.get_me(); text = message.text.replace(f"@{bot_user.username}", "").strip()
    pns = _split_pns(text)
    if not pns: return
    msg = await message.answer(f"⏳ Ищу {len(pns)} деталей..."); _user_last_pns[message.from_user.id] = pns
    for pn in pns: add_to_history(message.from_user.id, pn)
    results = await asyncio.to_thread(process_pns_sync, pns, api_key, force_update)
    if not results: await msg.edit_text("⚠️ Не найдено."); return
    out_text = format_text_result(results); kb = []
    if len(pns) == 1: kb.append([InlineKeyboardButton(text="🔄 Очистить кэш", callback_data=f"clear_{pns[0][:40]}")])
    kb.append([InlineKeyboardButton(text="📄 Выгрузить в Excel", callback_data="export_excel")])
    if len(out_text) > 4000:
        tmp_path = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name
        await asyncio.to_thread(write_xlsx, results, tmp_path)
        await message.answer_document(FSInputFile(tmp_path, filename="results.xlsx")); os.remove(tmp_path); await msg.delete()
    else: await msg.edit_text(out_text, reply_markup=InlineKeyboardMarkup(inline_keyboard=kb), link_preview_options=LinkPreviewOptions(is_disabled=True))

@dp.message(F.document)
async def handle_document(message: Message, bot: Bot, state: FSMContext) -> None:
    current_state = await state.get_state()
    if current_state == IPRegistryState.waiting_for_file.state:
        clear_ip_brands(); return await _import_categories_from_xlsx(message, bot, state, add_ip_brand)
    if current_state == DocsCategoryState.waiting_for_add.state:
        return await _import_categories_from_xlsx(message, bot, state, add_docs_category)
    if current_state == LabelingCategoryState.waiting_for_add.state:
        return await _import_categories_from_xlsx(message, bot, state, add_labeling_category)
    if current_state and current_state != UpdateState.waiting_for_data.state: return
    force_update = current_state == UpdateState.waiting_for_data.state
    if force_update: await state.clear()
    doc = message.document
    if not doc.file_name.lower().endswith('.xlsx'): await message.answer("Нужен .xlsx"); return
    msg = await message.answer("⏳ Обработка..."); tmp_in = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name
    await bot.download(doc, destination=tmp_in)
    parts, err = await asyncio.to_thread(extract_pns_from_excel_sync, tmp_in)
    if err == "NO_COLUMN":
        tpl = await asyncio.to_thread(create_template_excel); await message.answer_document(FSInputFile(tpl, filename="Template.xlsx"), caption="❌ Нет колонки Part no."); os.remove(tpl); os.remove(tmp_in); await msg.delete(); return
    elif err: os.remove(tmp_in); await msg.edit_text(f"❌ {err}"); return
    api_key = _get_api_key(); res_map = {}; total = len(parts); last_t = asyncio.get_event_loop().time()
    for idx, item in enumerate(parts):
        row = await asyncio.to_thread(_lookup_one, item['pn'], api_key, force_update=force_update)
        if row:
            if item['brand']: row['brand_from_excel'] = item['brand']
            res_map[item['pn'].lower()] = row
        now = asyncio.get_event_loop().time()
        if now - last_t > 1.0 or idx + 1 == total:
            await msg.edit_text(f"⏳ {idx+1}/{total}\n{get_progress_bar(idx+1, total)}"); last_t = now
    tmp_out = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name
    await asyncio.to_thread(enrich_xlsx, tmp_in, tmp_out, res_map)
    await message.answer_document(FSInputFile(tmp_out, filename=f"Enriched_{doc.file_name}")); os.remove(tmp_in); os.remove(tmp_out); await msg.delete()

async def _import_categories_from_xlsx(message: Message, bot: Bot, state: FSMContext, add_func: Callable[[str], bool]):
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name; await bot.download(message.document, destination=tmp)
    try:
        wb = load_workbook(tmp, data_only=True); count = sum(1 for row in wb.active.iter_rows(min_row=1, max_col=1) if row[0].value and add_func(str(row[0].value).strip()))
        await message.answer(f"✅ Добавлено: {count}")
    except Exception as e: await message.answer(f"❌ {e}")
    finally: os.remove(tmp); await state.clear()

@dp.callback_query(F.data == "export_excel")
async def callback_export_excel(callback: CallbackQuery):
    pns = _user_last_pns.get(callback.from_user.id)
    if not pns: await callback.answer("Устарело."); return
    await callback.message.edit_reply_markup(reply_markup=None); msg = await callback.message.answer("⏳ Формирую..."); api_key = _get_api_key()
    results = await asyncio.to_thread(process_pns_sync, pns, api_key); tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name
    await asyncio.to_thread(write_xlsx, results, tmp); await callback.message.answer_document(FSInputFile(tmp, filename="results.xlsx")); os.remove(tmp); await msg.delete(); await callback.answer()

@dp.callback_query(F.data.startswith("clear_"))
async def callback_clear_cache(callback: CallbackQuery):
    pn = callback.data.split("_", 1)[1]; delete_cached_part(pn); await callback.answer(f"Кэш {pn} очищен!")

class AuthMiddleware(BaseMiddleware):
    async def __call__(self, handler, event, data):
        if not ADMIN_USER_ID: return await handler(event, data)
        user = data.get('event_from_user')
        if not user or user.id == ADMIN_USER_ID or is_user_allowed(user.id): return await handler(event, data)
        if isinstance(event, types.Message) and event.chat.type == "private": await event.answer(f"⛔ Доступ ограничен. ID: {user.id}")

class CustomProxySession(AiohttpSession):
    def __init__(self, proxy_url: str, **kwargs):
        super().__init__(**kwargs)
        self.proxy_url = proxy_url

    async def create_session(self) -> aiohttp.ClientSession:
        if self._session is None or self._session.closed:
            if ProxyConnector:
                connector = ProxyConnector.from_url(self.proxy_url)
                self._session = aiohttp.ClientSession(connector=connector)
            else:
                logger.error("Библиотека aiohttp-socks не установлена! Прокси не будет работать.")
                self._session = aiohttp.ClientSession()
        return self._session

async def main() -> None:
    # Инициализируем сессию: через прокси или обычную
    if PROXY_URL:
        logger.info(f"Используется прокси: {PROXY_URL}")
        session = CustomProxySession(proxy_url=PROXY_URL)
    else:
        session = AiohttpSession()
    
    bot = Bot(
        token=TOKEN, 
        session=session,
        default=DefaultBotProperties(parse_mode=ParseMode.HTML)
    )
    
    dp.update.middleware(AuthMiddleware())
    
    # Команды ставим через try, чтобы не вешать запуск
    try:
        await bot.set_my_commands([
            BotCommand(command="menu", description="📱 Меню"), 
            BotCommand(command="history", description="🕒 История"), 
            BotCommand(command="stats", description="📊 База")
        ])
    except: pass

    logger.info("Бот запущен. Начинаю polling...")
    try:
        await dp.start_polling(bot)
    finally:
        await bot.session.close()

if __name__ == "__main__":
    # Фикс для Windows: SelectorEventLoop часто стабильнее на капризных сетевых драйверах
    if sys.platform == 'win32':
        asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        pass
    except Exception as e:
        logger.error(f"Критическая ошибка: {e}")